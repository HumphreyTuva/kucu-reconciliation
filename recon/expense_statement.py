"""Statement of actual expenses (PDF and Excel), laid out like the treasurer's expense report:
statement of expenses, notes, expenses by committee and month, and a basis of preparation."""
import calendar
import io
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.staticfiles import finders
from django.db.models import Q
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .budget import norm
from .defaults import OTHER_EXPENSES, statement_sort_key
from .models import BudgetLine, ManualEntry

ZERO = Decimal("0")
TITLE = "KENYATTA UNIVERSITY CHRISTIAN UNION (KUCU) - MAIN CAMPUS"
BRAND = "28309A"


def _prior_entered(pp):
    """Figures entered for the comparative year (e.g. from the audited accounts), by line."""
    qs = ManualEntry.objects.filter(Q(period__isnull=True, year=pp.end.year) | Q(period__end__range=(pp.start, pp.end)), section__startswith="out_")
    lines = {}
    for m in qs:
        label = "Special events" if m.section == "out_2_special_events" else m.label.strip()
        row = lines.setdefault(label.lower(), [label, ZERO])
        row[1] += m.amount
    return lines


def build(period):
    from .views import build_report  # imported here: views imports this module only inside a view

    r = build_report(period)
    prior = _prior_entered(period.prior())
    lines = []
    for x in r["expenses"]:
        label = x["label"]
        lines.append({"label": label, "note": 1 if label == "Special events" else 2 if label == OTHER_EXPENSES else None,
                      "cur": x["cur"], "prev": prior.pop(label.lower(), [label, ZERO])[1]})
    for label, amount in prior.values():  # lines that only exist in the comparative year, e.g. Depreciation
        lines.append({"label": label, "note": None, "cur": ZERO, "prev": amount})
    lines = [l for l in lines if l["cur"] or l["prev"]]
    lines.sort(key=lambda l: statement_sort_key(l["label"]))
    has_prev = any(l["prev"] for l in lines)
    total = {"cur": sum((l["cur"] for l in lines), ZERO), "prev": sum((l["prev"] for l in lines), ZERO)}

    months = period.months()
    budget = list(BudgetLine.objects.filter(period.month_filter(), actual__gt=0).select_related("category").order_by("id"))
    order, table, items, item_order = [], defaultdict(lambda: defaultdict(lambda: ZERO)), {}, []
    for b in budget:
        g = re.sub(r"\s+", " ", b.group).strip()
        if g not in order:
            order.append(g)
        table[g][(b.year, b.month)] += b.actual
        k = (g, norm(b.item))
        if k not in items:
            items[k] = {"group": g or "No committee", "item": re.sub(r"\s+", " ", b.item).strip(),
                        "line": b.category.name if b.category else f"{OTHER_EXPENSES} (no line chosen yet)", "months": defaultdict(lambda: ZERO)}
            item_order.append(k)
        items[k]["months"][(b.year, b.month)] += b.actual
    order.sort(key=lambda g: g == "")  # items outside any committee (e.g. Tuckshop) last
    committees = [{"name": g or "No committee (e.g. Tuckshop)", "months": [table[g][m] for m in months], "total": sum(table[g].values(), ZERO)} for g in order]
    month_totals = [sum((c["months"][i] for c in committees), ZERO) for i in range(len(months))]
    line_items = [{**items[k], "values": [items[k]["months"][m] for m in months], "total": sum(items[k]["months"].values(), ZERO)} for k in item_order]

    spent_months = {(b.year, b.month) for b in budget}
    empty = [m for m in months if m not in spent_months]
    other = r["note7"]
    unassigned = [o for o in other if o["unassigned"]]
    basis = [f"Period: {period.start:%d %B %Y} to {period.end:%d %B %Y}. Amounts are actual spending in Kshs; no budget figures are used."]
    if r["budget_months"]:
        basis.append(f"Expenses for {r['budget_months']} are taken line by line from the budget workbook. Bank debits in those months are not added again, "
                     "except bank charges, which are not in the workbook.")
    if empty:
        basis.append("No spending is recorded in the workbook for " + ", ".join(f"{calendar.month_abbr[m]} {y}" for y, m in empty) +
                     "; any bank debits in those months appear under Uncategorised expenses.")
    if other:
        basis.append(f"{len(other)} items have no report line of their own and are shown under Other expenses (Note 2), {sum((o['cur'] for o in other), ZERO):,.0f} in total."
                     + (f" {len(unassigned)} of them have not yet been assigned a line." if unassigned else ""))
    tuck = [o for o in other if "tuckshop" in o["label"].lower()]
    if tuck:
        basis.append(f"{tuck[0]['label']}, {tuck[0]['cur']:,.0f}, is included under Other expenses. Last year's accounts show the tuck shop as net project income, so please confirm this treatment.")
    if not any(l["label"].lower() in ("depreciation", "bank service charges") and l["cur"] for l in lines):
        basis.append("Depreciation and bank service charges are not recorded for this period, so the total is not like-for-like with a year that includes them.")
    if has_prev:
        basis.append(f"The {r['prev_label']} column shows the comparative figures entered for the previous year.")
    return {"period": period, "title": TITLE, "cur_label": r["cur_label"], "prev_label": r["prev_label"], "has_prev": has_prev, "lines": lines, "total": total,
            "note1": r["note6"][0], "note1_total": r["note6"][1], "note2": other, "note2_total": sum((o["cur"] for o in other), ZERO),
            "months": [f"{calendar.month_abbr[m]} {str(y)[2:]}" for y, m in months], "month_keys": months, "committees": committees, "month_totals": month_totals,
            "committee_total": sum(month_totals, ZERO), "line_items": line_items, "basis": basis}


# ------------------------------------------------------------------ Excel
def to_xlsx(d):
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=BRAND)
    num = '#,##0;(#,##0);"-"'
    top = Border(top=Side(style="thin"))

    def header(ws, row, values):
        for i, v in enumerate(values, 1):
            c = ws.cell(row, i, v)
            c.font, c.fill, c.alignment = head, fill, Alignment(horizontal="left" if i == 1 else "right", wrap_text=True)

    ws = wb.active
    ws.title = "Statement of Expenses"
    ws.append([d["title"]]); ws.append([f"STATEMENT OF ACTUAL EXPENSES - {d['period'].start:%B %Y} TO {d['period'].end:%B %Y}".upper()])
    ws.append(["Actual values only (no budget). Amounts in Kshs."]); ws.append([])
    ws["A1"].font = Font(bold=True, size=13); ws["A2"].font = Font(bold=True)
    cols = ["Expense", "Note", f"{d['cur_label']} Kshs"] + ([f"{d['prev_label']} Kshs"] if d["has_prev"] else [])
    header(ws, 5, cols)
    for l in d["lines"]:
        ws.append([l["label"], l["note"], float(l["cur"])] + ([float(l["prev"])] if d["has_prev"] else []))
    ws.append(["Total expenses", None, float(d["total"]["cur"])] + ([float(d["total"]["prev"])] if d["has_prev"] else []))
    for c in ws[ws.max_row]:
        c.font, c.border = Font(bold=True), top
    for row in ws.iter_rows(min_row=6, min_col=3):
        for c in row:
            c.number_format = num
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 7
    for c in "CD":
        ws.column_dimensions[c].width = 16
    ws.freeze_panes = "A6"

    li = wb.create_sheet("Line Items")
    header(li, 1, ["Committee", "Item", "Report line"] + d["months"] + ["Total"])
    for x in d["line_items"]:
        li.append([x["group"], x["item"], x["line"]] + [float(v) for v in x["values"]] + [float(x["total"])])
    li.append(["Total", None, None] + [float(v) for v in d["month_totals"]] + [float(d["committee_total"])])
    for c in li[li.max_row]:
        c.font, c.border = Font(bold=True), top
    for row in li.iter_rows(min_row=2, min_col=4):
        for c in row:
            c.number_format = num
    li.column_dimensions["A"].width = 18; li.column_dimensions["B"].width = 42; li.column_dimensions["C"].width = 36
    for i in range(4, 5 + len(d["months"])):
        li.column_dimensions[L(i)].width = 11
    li.freeze_panes = "D2"

    bc = wb.create_sheet("By Committee")
    header(bc, 1, ["Committee"] + d["months"] + ["Total"])
    for c_ in d["committees"]:
        bc.append([c_["name"]] + [float(v) for v in c_["months"]] + [float(c_["total"])])
    bc.append(["Monthly total"] + [float(v) for v in d["month_totals"]] + [float(d["committee_total"])])
    for c in bc[bc.max_row]:
        c.font, c.border = Font(bold=True), top
    for row in bc.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.number_format = num
    bc.column_dimensions["A"].width = 32
    for i in range(2, 3 + len(d["months"])):
        bc.column_dimensions[L(i)].width = 11
    bc.freeze_panes = "B2"

    nt = wb.create_sheet("Notes")
    nt.append(["Note 1 - Special events", f"{d['cur_label']} Kshs"])
    for x in d["note1"]:
        nt.append([x["label"], float(x["cur"])])
    nt.append(["Total", float(d["note1_total"])])
    nt.append([]); nt.append(["Note 2 - Other expenses (no report line of their own)", f"{d['cur_label']} Kshs"])
    for x in d["note2"]:
        nt.append([x["label"] + ("  (not yet assigned a line)" if x["unassigned"] else ""), float(x["cur"])])
    nt.append(["Total", float(d["note2_total"])])
    nt.append([]); nt.append(["Basis of preparation / items to confirm"])
    for i, b in enumerate(d["basis"], 1):
        nt.append([f"{i}. {b}"])
        nt.cell(nt.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        nt.row_dimensions[nt.max_row].height = 15 * (1 + len(b) // 110)
    for row in nt.iter_rows():
        if row[0].value and str(row[0].value).startswith(("Note", "Basis")):
            for c in row:
                c.font = Font(bold=True)
        if row[0].value == "Total":
            for c in row:
                c.font, c.border = Font(bold=True), top
        if len(row) > 1:
            row[1].number_format = num
    nt.column_dimensions["A"].width = 112; nt.column_dimensions["B"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF
def to_pdf(d):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, NextPageTemplate, PageBreak, PageTemplate,
                                    Paragraph, Spacer, Table, TableStyle)

    brand = colors.HexColor("#" + BRAND)
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12)
    bold = ParagraphStyle("bold", parent=base, fontName="Helvetica-Bold")
    small = ParagraphStyle("s", parent=base, fontSize=7, leading=9)
    h1 = ParagraphStyle("h1", parent=bold, fontSize=12, leading=15)
    h2 = ParagraphStyle("h2", parent=bold, fontSize=10, leading=13, textColor=brand, spaceBefore=10, spaceAfter=4)
    fmt = lambda v: "-" if not v else (f"({abs(v):,.0f})" if v < 0 else f"{v:,.0f}")
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, title="KUCU statement of actual expenses", author="KUCU Treasury")
    pw, ph = A4

    def footer(c, doc_):
        c.saveState(); c.setFont("Helvetica", 7); c.setFillColor(colors.grey)
        w, _ = c._pagesize
        c.drawString(15 * mm, 9 * mm, f"{d['title']} - statement of actual expenses")
        c.drawRightString(w - 15 * mm, 9 * mm, f"Page {doc_.page}")
        c.restoreState()

    doc.addPageTemplates([
        PageTemplate("portrait", [Frame(15 * mm, 15 * mm, pw - 30 * mm, ph - 28 * mm, id="p")], onPage=footer, pagesize=A4),
        PageTemplate("landscape", [Frame(12 * mm, 15 * mm, ph - 24 * mm, pw - 28 * mm, id="l")], onPage=footer, pagesize=landscape(A4))])

    def tbl(rows, widths, right_from=1, total_last=True, header=True, font=8.5):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        st = [("FONT", (0, 0), (-1, -1), "Helvetica", font), ("ALIGN", (right_from, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5), ("LINEBELOW", (0, 1), (-1, -2), 0.25, colors.HexColor("#e3e1d9"))]
        if header:
            st += [("BACKGROUND", (0, 0), (-1, 0), brand), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", font)]
        if total_last:
            st += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", font), ("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black), ("LINEBELOW", (0, -1), (-1, -1), 0.8, colors.black)]
        t.setStyle(TableStyle(st))
        return t

    story = []
    seal = finders.find("img/seal.png")
    head = [[Image(seal, 16 * mm, 16 * mm) if seal else "", [Paragraph(d["title"], h1), Paragraph(
        f"STATEMENT OF ACTUAL EXPENSES - {d['period'].start:%B %Y} TO {d['period'].end:%B %Y}".upper(), bold),
        Paragraph("Actual values only (no budget). Amounts in Kshs.", base)]]]
    ht = Table(head, colWidths=[20 * mm, 150 * mm]); ht.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    story += [ht, Spacer(1, 6)]
    cols = ["Expense", "Note", f"{d['cur_label']} Kshs"] + ([f"{d['prev_label']} Kshs"] if d["has_prev"] else [])
    rows = [cols] + [[Paragraph(l["label"], base), str(l["note"] or ""), fmt(l["cur"])] + ([fmt(l["prev"])] if d["has_prev"] else []) for l in d["lines"]]
    rows.append(["Total expenses", "", fmt(d["total"]["cur"])] + ([fmt(d["total"]["prev"])] if d["has_prev"] else []))
    story.append(tbl(rows, [95 * mm, 14 * mm, 32 * mm] + ([32 * mm] if d["has_prev"] else []), right_from=1))

    story.append(Paragraph("Note 1 - Special events", h2))
    story.append(tbl([["Event", f"{d['cur_label']} Kshs"]] + [[Paragraph(x["label"], base), fmt(x["cur"])] for x in d["note1"]] + [["Total", fmt(d["note1_total"])]], [141 * mm, 32 * mm]))
    story.append(Paragraph("Note 2 - Other expenses (items with no report line of their own)", h2))
    story.append(tbl([["Item", f"{d['cur_label']} Kshs"]] + [[Paragraph(x["label"] + (" *" if x["unassigned"] else ""), base), fmt(x["cur"])] for x in d["note2"]] + [["Total", fmt(d["note2_total"])]], [141 * mm, 32 * mm]))
    if any(x["unassigned"] for x in d["note2"]):
        story.append(Paragraph("* not yet assigned a report line of its own", small))

    story += [NextPageTemplate("landscape"), PageBreak(), Paragraph("Actual Expenses by Committee and Month", h2)]
    n = len(d["months"]); w = (ph - 24 * mm - 46 * mm - 20 * mm) / max(n, 1)
    crow = [["Committee"] + d["months"] + ["Total"]] + [[Paragraph(c["name"], small)] + [fmt(v) for v in c["months"]] + [fmt(c["total"])] for c in d["committees"]]
    crow.append(["Monthly total"] + [fmt(v) for v in d["month_totals"]] + [fmt(d["committee_total"])])
    story.append(tbl(crow, [46 * mm] + [w] * n + [20 * mm], font=7))
    story += [NextPageTemplate("portrait"), PageBreak(), Paragraph("Basis of preparation and items to confirm", h2)]
    for i, b in enumerate(d["basis"], 1):
        story += [Paragraph(f"{i}. {b}", base), Spacer(1, 4)]
    doc.build(story)
    return buf.getvalue()
