"""Reads the treasurer's budget workbook (monthly Actuals / Budget / Difference per expense line)."""
import calendar
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher

import openpyxl
from openpyxl.utils import get_column_letter as L

from .defaults import DEFAULTS, RETIRED
from .models import BudgetLine, Category, ItemMap
from .services import StatementError, _compile, find_category, to_dec

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
BANK_ONLY = {"Bank service charges"}  # comes from the bank statements, never from the budget workbook
STOP = re.compile(r"monthly totals?|total expenditure|grand total|^totals?$", re.I)


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()[:120]


def _amount(v):
    """A number, or a typed sum such as "9950+700"."""
    if v is None or v == "":
        return Decimal(0)
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "").replace(" ", "")
    if re.fullmatch(r"-?\d+(\.\d+)?([+\-]\d+(\.\d+)?)*", s):
        try:
            return sum((Decimal(p) for p in re.findall(r"[+\-]?\d+(?:\.\d+)?", s)), Decimal(0))
        except InvalidOperation:
            return Decimal(0)
    return Decimal(0)


def _is(cell, word):
    """Tolerant header match, so a typo such as "Actuas" still counts as "Actuals"."""
    return isinstance(cell, str) and SequenceMatcher(None, cell.strip().lower(), word).ratio() >= 0.8


def _month_of(label):
    """(year or None, month) for labels like JANUARY, October 2025, Oct-25 or the period code 202510."""
    if label is None:
        return None
    t = str(label).strip().lower()
    if t[:3] in MONTHS and not t.startswith("ytd"):
        m = MONTHS[t[:3]]
        ex = re.search(r"(20\d\d)", t) or re.search(r"[-/' ](\d\d)$", t)
        if ex:
            y = int(ex.group(1))
            return (y if y > 99 else 2000 + y), m
        return None, m
    m = re.fullmatch(r"(20\d\d)(\d\d)", t)
    if m and 1 <= int(m.group(2)) <= 12:
        return int(m.group(1)), int(m.group(2))
    return False


def parse_budget(upload, year=None, start=None, end=None):
    """Read every month column of the workbook. `year` is the year of the last month written without a year;
    earlier months that wrap past December belong to the year before. `start`/`end` (optional) limit the months."""
    try:
        wb = openpyxl.load_workbook(upload, data_only=True)
    except Exception:
        raise StatementError("This file could not be read as an Excel workbook.")
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), values_only=True))
        hr = next((i for i, r in enumerate(rows[:12]) if any(_is(c, "actuals") for c in r)
                   and any(_is(c, "budget") for c in r)), None)
        if hr is None or hr == 0:
            continue
        y = year or next((int(m) for m in re.findall(r"(20\d\d)", ws.title)), None) or date.today().year
        lab = next((c for r in rows[:hr + 1] for c, v in enumerate(r) if "expenditure" in str(v or "").lower()), 2)
        first_data = next((i for i in range(hr + 1, len(rows)) if str(rows[i][lab] or "").strip()
                           and not _is(rows[i][lab], "actuals") and not _is(rows[i][lab], "budget")), hr + 1)
        merged = {r.min_col - 1: r.max_col - 1 for r in ws.merged_cells.ranges if r.min_row < hr + 1}
        labels = sorted((c, v) for r in rows[:hr] for c, v in enumerate(r) if v not in (None, "") and c > lab)
        bounds = sorted({c for c, _ in labels} | set(merged))
        blocks, ignored, covered = [], [], set()
        for c, v in labels:
            nxt = next((x for x in bounds if x > c), c + 3)
            endc = merged.get(c, min(nxt - 1, c + 3))
            covered |= set(range(c, endc + 1))
            ym = _month_of(v)
            if not ym:
                ignored.append(f"column {L(c + 1)} ({str(v).strip()})")
                continue
            sub = [(cc, rows[r][cc]) for r in range(first_data) for cc in range(c, min(endc, len(rows[r]) - 1) + 1)]
            acol = next((cc for cc, val in sub if _is(val, "actuals")), c)
            bcol = next((cc for cc, val in sub if _is(val, "budget")), None)
            blocks.append({"col": acol, "bcol": bcol, "letter": L(c + 1), "label": str(v).strip(), "year": ym[0], "month": ym[1]})
        for r in rows[:first_data]:  # Actuals columns with no month heading at all
            for cc, val in enumerate(r):
                if _is(val, "actuals") and cc not in covered:
                    ignored.append(f"column {L(cc + 1)} (no month label)")
        if not blocks:
            continue
        seq = [b for b in blocks if b["year"] is None]  # months written without a year, in column order
        cur = y
        for k in range(len(seq) - 1, -1, -1):
            if k < len(seq) - 1 and seq[k]["month"] > seq[k + 1]["month"]:
                cur -= 1  # the year wrapped past December
            seq[k]["year"] = cur
        kept = []
        for b in blocks:
            first, last = date(b["year"], b["month"], 1), date(b["year"], b["month"], calendar.monthrange(b["year"], b["month"])[1])
            if (start and last < start) or (end and first > end):
                ignored.append(f"{b['label']} {b['year']} (outside the dates you set)")
            else:
                kept.append(b)
        blocks = kept
        lines, group, want_header = [], "", True
        for r in rows[first_data:]:
            name = re.sub(r"\s+", " ", str(r[lab] or "")).strip() if lab < len(r) else ""
            if not name or _is(name, "actuals") or _is(name, "budget"):
                continue
            if STOP.search(name):
                break
            if name.lower().startswith("subtotal"):
                want_header = True
                continue
            if want_header:
                carries_figures = any(
                    (b["col"] < len(r) and _amount(r[b["col"]])) or (b["bcol"] is not None and b["bcol"] < len(r) and _amount(r[b["bcol"]]))
                    for b in blocks)
                if not carries_figures:  # a department heading
                    group, want_header = name, False
                    continue
                group, want_header = "", False  # a lone line after a subtotal that has figures, e.g. "Tuckshop expenses"
            for bi, b in enumerate(blocks):
                a = _amount(r[b["col"]]) if b["col"] < len(r) else Decimal(0)
                bu = _amount(r[b["bcol"]]) if b["bcol"] is not None and b["bcol"] < len(r) else Decimal(0)
                if a or bu:
                    lines.append({"year": b["year"], "month": b["month"], "group": group, "item": name, "actual": a, "budget": bu, "_b": bi})
        # the same month can appear twice (e.g. OCTOBER and a 202510 period column): keep the fuller copy only
        tot = defaultdict(lambda: Decimal(0))
        for l in lines:
            tot[l["_b"]] += l["actual"] + l["budget"]
        keep = {}
        for bi, b in enumerate(blocks):
            k = (b["year"], b["month"])
            if k not in keep or tot[bi] > tot[keep[k]]:
                keep[k] = bi
        for bi, b in enumerate(blocks):
            if keep[(b["year"], b["month"])] != bi:
                what = "has no figures" if tot[bi] == 0 else f"repeats {calendar.month_name[b['month']]} {b['year']}"
                ignored.append(f"column {b['letter']} ({b['label']}) {what}")
        keep_ids = set(keep.values())
        lines = [{k: v for k, v in l.items() if k != "_b"} for l in lines if l["_b"] in keep_ids]
        blocks = [b for bi, b in enumerate(blocks) if bi in keep_ids]
        for b in blocks:
            mine = [l for l in lines if (l["year"], l["month"]) == (b["year"], b["month"])]
            b["actual"] = sum((l["actual"] for l in mine), Decimal(0))
            b["budget"] = sum((l["budget"] for l in mine), Decimal(0))
        ignored = list(dict.fromkeys(ignored))
        return {"sheet": ws.title, "blocks": blocks, "lines": lines, "ignored": ignored}
    raise StatementError("Could not find the Actuals / Budget columns in this workbook.")


def sync_defaults():
    """Make sure the standard report lines exist with their latest keywords (your own edits are kept), and retire
    lines that were replaced. Safe to run any number of times."""
    for i, (sec, name, kw) in enumerate(DEFAULTS):
        cat, made = Category.objects.get_or_create(name=name, section=sec, defaults={"keywords": kw, "order": 10 + i})
        if not made:
            have = [k.strip() for k in cat.keywords.split(",") if k.strip()]
            new = [k for k in (x.strip() for x in kw.split(",")) if k and k.lower() not in [h.lower() for h in have]]
            if new:
                cat.keywords = ",".join(have + new)[:300]
                cat.save(update_fields=["keywords"])
    for sec, name in RETIRED:
        cat = Category.objects.filter(name=name, section=sec).first()
        if cat:
            if cat.keywords:
                cat.keywords = ""
                cat.save(update_fields=["keywords"])
            BudgetLine.objects.filter(category=cat).update(category=None)  # re-mapped below
            ItemMap.objects.filter(category=cat).delete()


def import_budget(parsed):
    """Store every month that has spending or a planned budget. A month uploaded again replaces the earlier copy.
    Returns the months with actual spending, the months with only a planned budget, and the empty ones."""
    sync_defaults()
    expense = _compile([c for c in Category.objects.filter(is_active=True) if not c.is_income and c.name not in BANK_ONLY])
    saved = {i.key: i.category for i in ItemMap.objects.select_related("category")}
    done = {"actual": [], "planned": [], "empty": []}
    keep = set()
    for b in parsed["blocks"]:
        ym = (b["year"], b["month"])
        kind = "actual" if b["actual"] > 0 else "planned" if b["budget"] > 0 else "empty"
        if ym not in done[kind] and all(ym not in done[k] for k in done):
            done[kind].append(ym)
        if kind != "empty":
            keep.add(ym)
    cache, out = {}, []
    for l in parsed["lines"]:
        if (l["year"], l["month"]) not in keep:
            continue
        cat = saved.get(norm(l["item"]))
        if cat is None:
            cat, _ = find_category(l["item"], expense, True, cache)
        if cat is not None and cat.name in BANK_ONLY:
            cat = None
        out.append(BudgetLine(category=cat, **l))
    # the file is the latest word on every month it covers, so months it now shows as empty are cleared too
    for y, m in {(b["year"], b["month"]) for b in parsed["blocks"]}:
        BudgetLine.objects.filter(year=y, month=m).delete()
    BudgetLine.objects.bulk_create(out, batch_size=500)
    for k in done:
        done[k].sort()
    return done


def automap_unmapped():
    """Apply the saved choices and the keyword rules to every budget line that has no report line yet."""
    sync_defaults()
    expense = _compile([c for c in Category.objects.filter(is_active=True) if not c.is_income and c.name not in BANK_ONLY])
    saved = {i.key: i.category for i in ItemMap.objects.select_related("category")}
    cache, changed = {}, []
    for l in BudgetLine.objects.filter(category__isnull=True):
        cat = saved.get(norm(l.item)) or find_category(l.item, expense, True, cache)[0]
        if cat:
            l.category = cat
            changed.append(l)
    BudgetLine.objects.bulk_update(changed, ["category"], batch_size=500)
    return len(changed)
