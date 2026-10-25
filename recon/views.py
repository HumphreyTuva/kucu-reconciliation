import calendar
import io
from collections import defaultdict
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib import messages
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Max, Min, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from .budget import BANK_ONLY
from .forms import CategoryFormSet, UploadForm
from .models import SECTIONS, BudgetLine, Category, FinancialYear, ItemMap, ManualEntry, Statement, Transaction
from .services import StatementError, categorise, import_statement, parse_statement

SEC_LABEL = dict(SECTIONS)
ZERO = Decimal("0")


def groups(cats):
    out = defaultdict(list)
    for c in cats:
        out[c.section].append(c)
    return [(SEC_LABEL[k], v) for k, v in sorted(out.items(), key=lambda kv: (kv[0].startswith("ex_"), kv[0]))]


# ---------------------------------------------------------------- statements
@login_required
def home(request):
    form = UploadForm()
    if request.method == "POST":
        form = UploadForm(request.POST, request.FILES)
        if form.is_valid():
            made = []
            for f in form.cleaned_data["file"]:
                try:
                    rows, opening, closing = parse_statement(f, form.cleaned_data.get("pdf_password", ""))
                except StatementError as e:
                    messages.error(request, f"{f.name}: {e}")
                    continue
                st = Statement.objects.create(name=f.name, opening=opening, closing=closing)
                import_statement(st, rows)
                made.append(st)
                if opening is not None and closing is not None:
                    gap = opening + sum(r["credit"] for r in rows) - sum(r["debit"] for r in rows) - closing
                    if gap != 0:
                        messages.warning(request, f"{f.name}: the totals do not balance (off by {gap:,.2f}), so some transactions may have been misread. "
                                                  "If this is a PDF, try the Excel version of the statement and compare.")
                seen = set(Transaction.objects.exclude(statement=st).values_list("date", "description", "debit", "credit"))
                dups = sum(1 for t in st.transactions.all() if (t.date, t.description, t.debit, t.credit) in seen)
                if dups:
                    messages.warning(request, f"{f.name}: {dups} of {len(rows)} transactions already exist in another statement. "
                                              "The yearly report would count them twice, so delete one of the statements if this was a repeat upload.")
            if len(made) == 1:
                return redirect("statement", pk=made[0].pk)
            if made:
                messages.success(request, f"{len(made)} statements reconciled: " + ", ".join(m.name for m in made) + ". Open any of them below, or see the Annual report for the combined totals.")
                return redirect("home")
    statements = Statement.objects.annotate(first=Min("transactions__date"), last=Max("transactions__date")).order_by("-first", "-uploaded_at")[:36]
    return render(request, "recon/home.html", {"form": form, "statements": statements})


def _lines(cats, agg, key, name):
    out = [{"key": str(c.id), "name": c.name + (" (not counted)" if c.is_excluded else ""), "n": agg[c.id]["n"], "total": agg[c.id]["total"]}
           for c in cats if c.id in agg]
    if None in agg:
        out.append({"key": key, "name": name, "n": agg[None]["n"], "total": agg[None]["total"], "warn": True})
    tot = sum((l["total"] for l in out), ZERO)
    for l in out:
        l["pct"] = float(l["total"] / tot * 100) if tot else 0
    return out, tot


def build_summary(st):
    cats = list(Category.objects.all())
    agg = lambda f, fld: {r["category"]: r for r in st.transactions.filter(**{f + "__gt": 0}).values("category")
                          .annotate(n=Count("id"), total=Sum(fld))}
    income, cr = _lines(cats, agg("credit", "credit"), "nocr", "Uncategorised credits")
    expense, db = _lines(cats, agg("debit", "debit"), "nodb", "Uncategorised debits")
    return income, cr, expense, db


@login_required
def statement(request, pk):
    st = get_object_or_404(Statement, pk=pk)
    income, cr, expense, db = build_summary(st)
    check = st.opening + cr - db - st.closing if st.opening is not None and st.closing is not None else None
    flt = request.GET.get("cat", "")
    qs = st.transactions.select_related("category")
    if flt == "nocr":
        qs = qs.filter(credit__gt=0, category__isnull=True)
    elif flt == "nodb":
        qs = qs.filter(debit__gt=0, category__isnull=True)
    elif flt.isdigit():
        qs = qs.filter(category_id=int(flt))
    from django.db.models import Max as _Max, Min as _Min
    span = st.transactions.aggregate(a=_Min("date"), b=_Max("date"))
    page = Paginator(qs, 100).get_page(request.GET.get("page"))
    active = list(Category.objects.filter(is_active=True))
    return render(request, "recon/statement.html", {
        "st": st, "income": income, "expense": expense, "total_credit": cr, "total_debit": db,
        "net": cr - db, "check": check, "txs": page, "flt": flt,
        "income_groups": groups([c for c in active if c.is_income or c.is_excluded]),
        "expense_groups": groups([c for c in active if not c.is_income]),
        "start": span["a"], "end": span["b"],
    })


@login_required
@require_POST
def set_category(request, pk):
    t = get_object_or_404(Transaction, pk=pk)
    cid = request.POST.get("category")
    t.category = Category.objects.filter(pk=cid).first() if cid else None
    t.manual, t.note = True, ""
    t.save(update_fields=["category", "manual", "note"])
    return redirect(f"/statement/{t.statement_id}/?cat={request.POST.get('cat', '')}&page={request.POST.get('page', '1')}#tx")


@login_required
@require_POST
def recategorise(request, pk):
    st = get_object_or_404(Statement, pk=pk)
    if request.POST.get("reset_manual"):
        st.transactions.update(manual=False)
    categorise(st)
    messages.success(request, "Category rules re-applied.")
    return redirect("statement", pk=pk)


@login_required
@require_POST
def delete_statement(request, pk):
    get_object_or_404(Statement, pk=pk).delete()
    return redirect("home")


def _xlsx_response(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp


@login_required
def export(request, pk):
    st = get_object_or_404(Statement, pk=pk)
    income, cr, expense, db = build_summary(st)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"KUCU reconciliation: {st.name}"]); ws.append([])
    ws.append(["Income (credits)", "Transactions", "Amount (KES)"])
    for l in income:
        ws.append([l["name"], l["n"], float(l["total"])])
    ws.append(["Total credits", sum(l["n"] for l in income), float(cr)]); ws.append([])
    ws.append(["Expenses (debits)", "Transactions", "Amount (KES)"])
    for l in expense:
        ws.append([l["name"], l["n"], float(l["total"])])
    ws.append(["Total debits", sum(l["n"] for l in expense), float(db)])
    if st.opening is not None:
        ws.append([]); ws.append(["Opening balance", None, float(st.opening)])
    if st.closing is not None:
        ws.append(["Closing balance", None, float(st.closing)])
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["C"].width = 18
    for r in ws.iter_rows(min_col=3, max_col=3):
        r[0].number_format = "#,##0.00"
    tx = wb.create_sheet("Transactions")
    tx.append(["Date", "Details", "Debit", "Credit", "Category"])
    for t in st.transactions.select_related("category"):
        tx.append([t.date, t.description, float(t.debit) or None, float(t.credit) or None,
                   t.category.name if t.category else "Uncategorised"])
    tx.column_dimensions["B"].width = 80
    return _xlsx_response(wb, f"reconciliation-{st.pk}")


@login_required
def categories(request):
    qs = Category.objects.order_by("section", "order", "name")
    if request.method == "POST":
        fs = CategoryFormSet(request.POST, queryset=qs)
        if fs.is_valid():
            fs.save()
            messages.success(request, "Categories saved. Open a statement and press Re-run rules to apply them.")
            return redirect("categories")
    else:
        fs = CategoryFormSet(queryset=qs)
    return render(request, "recon/categories.html", {"formset": fs})


# -------------------------------------------------------------------- periods
from dataclasses import dataclass
from datetime import datetime as _dt


def _shift(d, years):
    try:
        return d.replace(year=d.year + years)
    except ValueError:  # 29 February
        return d.replace(year=d.year + years, day=28)


@dataclass
class Period:
    id: str
    name: str
    start: date
    end: date
    defined: bool

    def prior(self):
        return Period("prior", "", _shift(self.start, -1), _shift(self.end, -1), False)

    def months(self):
        out, y, m = [], self.start.year, self.start.month
        while (y, m) <= (self.end.year, self.end.month):
            out.append((y, m))
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
        return out

    @property
    def is_year(self):
        return self.defined and 330 <= (self.end - self.start).days + 1 <= 400

    @property
    def heading(self):
        if self.is_year:
            return f"FOR THE YEAR ENDED {self.end:%d %B %Y}".upper()
        return f"FOR THE PERIOD {self.start:%d %B %Y} TO {self.end:%d %B %Y}".upper()

    @property
    def cur_label(self):
        return str(self.end.year) if self.is_year else f"{self.start:%b %Y} - {self.end:%b %Y}"

    @property
    def prev_label(self):
        p = self.prior()
        return str(p.end.year) if self.is_year else f"{p.start:%b %Y} - {p.end:%b %Y}"

    def month_filter(self):
        sy, sm, ey, em = self.start.year, self.start.month, self.end.year, self.end.month
        return (Q(year__gt=sy) | Q(year=sy, month__gte=sm)) & (Q(year__lt=ey) | Q(year=ey, month__lte=em))


def data_bounds(budget_only=False):
    """First and last day covered by what has been uploaded."""
    months = list(BudgetLine.objects.values_list("year", "month").distinct())
    starts = [date(y, m, 1) for y, m in months]
    ends = [date(y, m, calendar.monthrange(y, m)[1]) for y, m in months]
    if not budget_only:
        t = Transaction.objects.exclude(date=None).aggregate(a=Min("date"), b=Max("date"))
        starts += [x for x in (t["a"],) if x]
        ends += [x for x in (t["b"],) if x]
    if not starts:
        return date.today(), date.today()
    return min(starts), max(ends)


def resolve_period(request, budget_only=False):
    """The financial year chosen with ?p=<id>, else the latest one you set up, else everything uploaded."""
    fys = list(FinancialYear.objects.all())
    p = request.GET.get("p", "")
    chosen = next((f for f in fys if str(f.id) == p), None) if p and p != "all" else None
    if chosen is None and not p and fys:
        chosen = fys[0]
    if chosen:
        return Period(str(chosen.id), chosen.label, chosen.start, chosen.end, True), fys
    start, end = data_bounds(budget_only)
    return Period("all", "All data", start, end, False), fys


def _fy_name(start, end):
    if 330 <= (end - start).days + 1 <= 400:
        return str(end.year) if start.year == end.year else f"{start.year}/{str(end.year)[-2:]}"
    return f"{start:%b %Y} - {end:%b %Y}"


def _parse_date(text):
    try:
        return _dt.strptime((text or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


@login_required
def periods(request):
    if request.method == "POST":
        if request.POST.get("action") == "delete":
            FinancialYear.objects.filter(pk=request.POST.get("id")).delete()
            messages.success(request, "Financial year removed. Nothing else was deleted.")
        else:
            start, end = _parse_date(request.POST.get("start")), _parse_date(request.POST.get("end"))
            if not start or not end or end < start:
                messages.error(request, "Enter a start date and an end date that comes after it.")
            else:
                FinancialYear.objects.get_or_create(start=start, end=end, defaults={"name": request.POST.get("name", "").strip()[:80] or _fy_name(start, end)})
                messages.success(request, "Financial year added.")
        return redirect("periods")
    return render(request, "recon/periods.html", {"fys": FinancialYear.objects.all(), "bounds": data_bounds()})


# -------------------------------------------------------------------- report
def year_data(period, prior=False):
    start, end = period.start, period.end
    qs = Transaction.objects.filter(date__range=(start, end))
    # Months covered by the budget workbook take their expenses from its itemised actuals, not from bank debits
    # (the same spending would otherwise be counted twice).
    bud = list(BudgetLine.objects.filter(period.month_filter()).select_related("category"))
    spent = defaultdict(lambda: ZERO)
    for b in bud:
        spent[(b.year, b.month)] += b.actual
    bud_months = sorted(k for k, v in spent.items() if v > 0)  # planned-only months do not replace bank debits
    in_budget_month = Q()
    for yy, mm in bud_months:
        in_budget_month |= Q(date__year=yy, date__month=mm)
    # bank fees are never itemised in the workbook, so they still come from the statements
    qs_out = qs.filter(~in_budget_month | Q(category__name__in=BANK_ONLY)) if bud_months else qs
    amt, uncat, unmapped = defaultdict(lambda: ZERO), 0, 0
    for fld, base, prefix, fallback in (("credit", qs, "in_", ("in_3_other", "Other incomes (uncategorised)")),
                                        ("debit", qs_out, "out_", ("out_1_expenses", "Uncategorised expenses"))):
        for r in base.filter(**{fld + "__gt": 0}).values("category__section", "category__name").annotate(t=Sum(fld), n=Count("id")):
            sec = r["category__section"]
            if sec and sec.startswith("ex_"):
                continue  # transfers between own accounts are not income or expense
            if sec and sec.startswith(prefix):
                amt[(sec, r["category__name"])] += r["t"]
            else:
                amt[fallback] += r["t"]; uncat += r["n"]
    for b in bud:
        if not b.actual:
            continue
        cat = b.category
        if cat and cat.section.startswith("ex_"):
            continue
        if cat and cat.section.startswith("out_"):
            amt[(cat.section, cat.name)] += b.actual
        else:
            amt[("out_1_expenses", "Other expenses (not yet mapped)")] += b.actual
            unmapped += 1
    if prior:
        mq = Q(period__isnull=True, year=end.year) | Q(period__end__range=(start, end))
    elif period.defined:
        mq = Q(period_id=int(period.id)) | Q(period__isnull=True, year=end.year)
    else:
        mq = Q()  # all data: every manual entry
    manual = list(ManualEntry.objects.filter(mq))
    for m in manual:
        amt[(m.section, m.label)] += m.amount
    return amt, uncat, qs.exists() or bool(manual) or bool(bud), bud_months, unmapped


def build_report(period):
    cur, uncat, has_cur, bud_months, unmapped = year_data(period)
    if period.defined:
        prev, _, has_prev, _, _ = year_data(period.prior(), prior=True)
    else:
        prev, has_prev = defaultdict(lambda: ZERO), False

    def rows(sec):
        names = [c.name for c in Category.objects.filter(section=sec)]
        extra = sorted({l for (s, l) in list(cur) + list(prev) if s == sec and l not in names})
        out = [{"label": n, "cur": cur.get((sec, n), ZERO), "prev": prev.get((sec, n), ZERO)} for n in names + extra]
        out = [r for r in out if r["cur"] or r["prev"]]
        return out, sum((r["cur"] for r in out), ZERO), sum((r["prev"] for r in out), ZERO)

    S = {k: rows(k) for k, _ in SECTIONS}
    row = lambda label, k, note=None: {"label": label, "note": note, "cur": S[k][1], "prev": S[k][2]}
    income = [row("Tithes and Offering", "in_1_tithes_offering", 5), row("Income from projects", "in_2_projects", 5), row("Other income", "in_3_other", 5)]
    expenses = [{"label": r["label"], "cur": r["cur"], "prev": r["prev"]} for r in S["out_1_expenses"][0]]
    expenses.append(row("Special events", "out_2_special_events", 6))
    tot = lambda rs, f: sum((r[f] for r in rs), ZERO)
    ti, te = {f: tot(income, f) for f in ("cur", "prev")}, {f: tot(expenses, f) for f in ("cur", "prev")}
    return {
        "heading": period.heading, "cur_label": period.cur_label, "prev_label": period.prev_label,
        "income": income, "expenses": expenses, "total_income": ti, "total_expenses": te,
        "surplus": {f: ti[f] - te[f] for f in ("cur", "prev")},
        "note5": [(title, S[k]) for title, k in (("Tithes and Offering", "in_1_tithes_offering"),
                                                  ("Net proceeds from projects", "in_2_projects"), ("Other income", "in_3_other"))],
        "note6": S["out_2_special_events"], "uncat": uncat, "has_data": has_cur, "has_prev": has_prev,
        "budget_months": ", ".join(f"{calendar.month_abbr[m]} {yy}" for yy, m in bud_months), "unmapped": unmapped,
    }


@login_required
def report(request):
    period, fys = resolve_period(request)
    has_any = Transaction.objects.exists() or BudgetLine.objects.exists() or ManualEntry.objects.exists()
    ctx = {"fys": fys, "period": period, "p": period.id, "has_any": has_any}
    if has_any:
        ctx.update(build_report(period))
    return render(request, "recon/report.html", ctx)


@login_required
def report_export(request):
    period, _ = resolve_period(request)
    r = build_report(period)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Income statement"
    for t in ("KENYATTA UNIVERSITY CHRISTIAN UNION (KUCU) - MAIN CAMPUS", "REPORT AND FINANCIAL STATEMENTS",
              r["heading"], "STATEMENT OF COMPREHENSIVE INCOME", ""):
        ws.append([t])
    head = ["", "Notes", f"{r['cur_label']} Kshs"] + ([f"{r['prev_label']} Kshs"] if r["has_prev"] else [])
    ws.append(head)

    def put(label, note, c, p, bold=False):
        ws.append([label, note, float(c)] + ([float(p)] if r["has_prev"] else []))
        for cell in ws[ws.max_row]:
            cell.font = openpyxl.styles.Font(bold=bold)
        for cell in ws[ws.max_row][2:]:
            cell.number_format = '#,##0;(#,##0);"-"'

    ws.append(["INCOME"])
    for x in r["income"]:
        put(x["label"], x["note"], x["cur"], x["prev"])
    put("Total Income", None, r["total_income"]["cur"], r["total_income"]["prev"], True)
    ws.append([]); ws.append(["EXPENSES"])
    for x in r["expenses"]:
        put(x["label"], x.get("note"), x["cur"], x["prev"])
    put("Total expenses", None, r["total_expenses"]["cur"], r["total_expenses"]["prev"], True)
    ws.append([])
    put("Surplus/(Deficit) for the period", None, r["surplus"]["cur"], r["surplus"]["prev"], True)
    ws.column_dimensions["A"].width = 52
    for c in "CD":
        ws.column_dimensions[c].width = 16
    notes = wb.create_sheet("Notes")

    def nrow(label, cur, prev):
        notes.append([label, "", float(cur)] + ([float(prev)] if r["has_prev"] else []))

    notes.append(["NOTE 5 - INCOME", "", f"{r['cur_label']} Kshs"] + ([f"{r['prev_label']} Kshs"] if r["has_prev"] else []))
    for title, (rows_, tc, tp) in r["note5"]:
        notes.append([title])
        for x in rows_:
            nrow(x["label"], x["cur"], x["prev"])
        nrow(f"Total {title.lower()}", tc, tp)
    notes.append([]); notes.append(["NOTE 6 - SPECIAL EVENTS"])
    for x in r["note6"][0]:
        nrow(x["label"], x["cur"], x["prev"])
    nrow("Total special events", r["note6"][1], r["note6"][2])
    notes.column_dimensions["A"].width = 52
    for rw in notes.iter_rows(min_col=3, max_col=4):
        for cell in rw:
            cell.number_format = '#,##0;(#,##0);"-"'
    return _xlsx_response(wb, f"KUCU-financial-report-{period.end:%Y%m%d}")


# ----------------------------------------------------------------------- PWA
import json

from django.templatetags.static import static
from django.views.decorators.cache import never_cache


def manifest(request):
    data = {
        "name": "KUCU Treasury Reconciliation", "short_name": "KUCU Treasury",
        "description": "Reconcile bank statements and build the annual financial report.",
        "start_url": "/", "scope": "/", "display": "standalone",
        "background_color": "#ffffff", "theme_color": "#28309a",
        "icons": [
            {"src": static("img/icon-192.png") + "?v=2", "sizes": "192x192", "type": "image/png"},
            {"src": static("img/icon-512.png") + "?v=2", "sizes": "512x512", "type": "image/png"},
            {"src": static("img/icon-maskable-512.png") + "?v=2", "sizes": "512x512", "type": "image/png", "purpose": "maskable"},
        ],
    }
    return HttpResponse(json.dumps(data), content_type="application/manifest+json")


@never_cache
def service_worker(request):
    js = render(request, "sw.js", {"css": static("css/app.css"), "logo": static("img/logo-wide.png"),
                                   "icon": static("img/icon-192.png") + "?v=2"}).content
    resp = HttpResponse(js, content_type="application/javascript")
    resp["Service-Worker-Allowed"] = "/"
    return resp


def offline(request):
    return render(request, "offline.html")


# ------------------------------------------------------------------- support
def support(request):
    """Public page, so people who cannot log in can still reach support."""
    return render(request, "recon/support.html", {"phone": "+254113631232", "wa": "254113631232", "email": "tuvalabs@gmail.com"})


# -------------------------------------------------------------------- budget
from .budget import automap_unmapped, import_budget, norm, parse_budget


def _names(months):
    return ", ".join(f"{calendar.month_abbr[m]} {y}" for y, m in months)


@login_required
def budget(request):
    if request.method == "POST":
        f = request.FILES.get("file")
        year = request.POST.get("year", "").strip()
        start, end = _parse_date(request.POST.get("start")), _parse_date(request.POST.get("end"))
        if not f:
            messages.error(request, "Choose the budget workbook first.")
        elif (request.POST.get("start") or request.POST.get("end")) and not (start and end and end >= start):
            messages.error(request, "Set both the start and the end date of the financial year, or leave both empty to read everything in the workbook.")
        else:
            try:
                parsed = parse_budget(f, int(year) if year.isdigit() else None, start, end)
                if request.POST.get("fresh") == "1":  # start fresh: remove everything imported before
                    BudgetLine.objects.all().delete()
                done = import_budget(parsed)
            except StatementError as e:
                messages.error(request, f"{f.name}: {e}")
            else:
                if done["actual"] or done["planned"]:
                    msg = f"Imported {len(parsed['lines'])} lines from sheet \u201c{parsed['sheet']}\u201d. "
                    if done["actual"]:
                        msg += f"Spending found for: {_names(done['actual'])}. "
                    if done["planned"]:
                        msg += f"Planned budget only: {_names(done['planned'])}."
                    messages.success(request, msg.strip())
                else:
                    messages.warning(request, "This workbook has no spending or budget figures yet, so nothing was imported.")
                if done["empty"]:
                    messages.info(request, f"Nothing recorded in the workbook for: {_names(done['empty'])}.")
                if parsed["ignored"]:
                    messages.info(request, "Skipped: " + "; ".join(parsed["ignored"]) + ".")
                if start and end:
                    fy, made = FinancialYear.objects.get_or_create(start=start, end=end, defaults={"name": _fy_name(start, end)})
                    if made:
                        messages.info(request, f"Financial year {fy.label} ({start:%d %b %Y} to {end:%d %b %Y}) was saved. Change it under Financial years.")
                    return redirect(f"/budget/?p={fy.id}")
                return redirect("/budget/?p=all")
    period, fys = resolve_period(request, budget_only=True)
    rows, tot = [], {"spent": ZERO, "budget_all": ZERO, "spent_cmp": ZERO, "budget_cmp": ZERO, "planned": ZERO}
    has_lines = BudgetLine.objects.exists()
    if has_lines:
        for yy, mm in period.months():
            a = BudgetLine.objects.filter(year=yy, month=mm).aggregate(a=Sum("actual"), b=Sum("budget"), n=Count("id"))
            has, act = a["n"] > 0, (a["a"] or ZERO) > 0
            bank = Transaction.objects.filter(date__year=yy, date__month=mm, debit__gt=0).aggregate(t=Sum("debit"))["t"]
            budget_m, actual_m = a["b"] or ZERO, a["a"] or ZERO
            rows.append({"year": yy, "month": mm, "name": f"{calendar.month_name[mm]} {yy}", "status": "actual" if act else "planned" if has else "none",
                         "actual": actual_m, "budget": budget_m, "has_budget": budget_m > 0, "diff": budget_m - actual_m, "bank": bank})
            tot["budget_all"] += budget_m
            if has and not act:
                tot["planned"] += budget_m  # budgeted for months that have no spending yet
            if act:
                tot["spent"] += actual_m
                if budget_m > 0:  # compare only months that have both
                    tot["spent_cmp"] += actual_m
                    tot["budget_cmp"] += budget_m
    unmapped = BudgetLine.objects.filter(category__isnull=True).values("item").distinct().count()
    return render(request, "recon/budget.html", {
        "period": period, "p": period.id, "fys": fys, "rows": rows, "tot": tot, "has_lines": has_lines,
        "diff": tot["budget_cmp"] - tot["spent_cmp"], "remaining": tot["planned"],
        "months_spent": sum(1 for r in rows if r["status"] == "actual"), "unmapped": unmapped})


@login_required
@require_POST
def budget_delete(request):
    scope, p = request.POST.get("scope"), request.POST.get("p", "")
    try:
        if scope == "month":
            n = BudgetLine.objects.filter(year=int(request.POST["year"]), month=int(request.POST["month"])).delete()[0]
            messages.success(request, f"Deleted {n} budget lines for {calendar.month_name[int(request.POST['month'])]} {request.POST['year']}.")
        elif scope == "period":
            fy = FinancialYear.objects.filter(pk=p).first() if p.isdigit() else None
            if fy:
                n = BudgetLine.objects.filter(Period(str(fy.id), "", fy.start, fy.end, True).month_filter()).delete()[0]
                messages.success(request, f"Deleted {n} budget lines in {fy.label}. Your report-line choices are kept, so a fresh import maps itself.")
            else:
                n = BudgetLine.objects.all().delete()[0]
                messages.success(request, f"Deleted all {n} budget lines. Your report-line choices are kept, so a fresh import maps itself.")
    except (KeyError, ValueError):
        messages.error(request, "Nothing was deleted.")
    return redirect(f"/budget/?p={p}" if p else "budget")


@login_required
def budget_month(request, year, month):
    lines = list(BudgetLine.objects.filter(year=year, month=month).select_related("category"))
    if not lines:
        messages.error(request, "No budget data for that month.")
        return redirect("budget")
    groups_, order = {}, []
    for l in lines:
        if l.group not in groups_:
            groups_[l.group] = []
            order.append(l.group)
        groups_[l.group].append({"item": l.item, "actual": l.actual, "budget": l.budget, "diff": l.budget - l.actual, "cat": l.category})
    out = [{"name": g or "Other", "lines": groups_[g], "actual": sum((x["actual"] for x in groups_[g]), ZERO),
            "budget": sum((x["budget"] for x in groups_[g]), ZERO)} for g in order]
    for g in out:
        g["diff"] = g["budget"] - g["actual"]
    ta, tb = sum((g["actual"] for g in out), ZERO), sum((g["budget"] for g in out), ZERO)
    bank = Transaction.objects.filter(date__year=year, date__month=month, debit__gt=0).aggregate(t=Sum("debit"))["t"]
    return render(request, "recon/budget_month.html", {"title": f"{calendar.month_name[month]} {year}", "groups": out, "actual": ta,
                                                        "budget": tb, "diff": tb - ta, "bank": bank, "p": request.GET.get("p", ""), "has_budget": tb > 0})


@login_required
def budget_mapping(request):
    cats = [c for c in Category.objects.filter(is_active=True) if c.section.startswith("out_") or c.is_excluded]
    if request.method == "POST":
        names = list(BudgetLine.objects.values_list("item", flat=True).distinct())
        by_key = defaultdict(list)
        for n in names:
            by_key[norm(n)].append(n)
        for key, value in request.POST.items():
            if not key.startswith("map_"):
                continue
            k = key[4:]
            if value == "__new__":  # a report line typed in on this page
                name = request.POST.get(f"newname_{k}", "").strip()[:80]
                sec = request.POST.get(f"newsec_{k}", "out_1_expenses")
                if not name:
                    continue
                if sec not in ("out_1_expenses", "out_2_special_events"):
                    sec = "out_1_expenses"
                cat = Category.objects.filter(name__iexact=name, section=sec).first() or \
                    Category.objects.create(name=name, section=sec, keywords="", order=200)
            else:
                cat = Category.objects.filter(pk=value).first() if value else None
            if cat:
                ItemMap.objects.update_or_create(key=k, defaults={"category": cat})
            else:
                ItemMap.objects.filter(key=k).delete()
            BudgetLine.objects.filter(item__in=by_key.get(k, [])).update(category=cat)
        messages.success(request, "Mapping saved. The annual report now uses it.")
        return redirect(request.get_full_path())  # stay on the same filter
    agg = {}
    for l in BudgetLine.objects.select_related("category"):
        k = norm(l.item)
        a = agg.setdefault(k, {"key": k, "item": l.item, "groups": set(), "actual": ZERO, "budget": ZERO, "cat": None, "unmapped": False})
        a["actual"] += l.actual
        a["budget"] += l.budget
        a["groups"].add(l.group)
        if l.category_id:
            a["cat"] = a["cat"] or l.category_id
        else:
            a["unmapped"] = True
    items = sorted(agg.values(), key=lambda a: (-a["actual"], -a["budget"]))
    counts = {"all": len(items), "unmapped": sum(1 for a in items if a["unmapped"])}
    counts["mapped"] = counts["all"] - counts["unmapped"]
    show = request.GET.get("show", "all")
    if show not in ("all", "unmapped", "mapped"):
        show = "all"
    q = request.GET.get("q", "").strip().lower()
    if show == "unmapped":
        items = [a for a in items if a["unmapped"]]
    elif show == "mapped":
        items = [a for a in items if not a["unmapped"]]
    for a in items:
        a["groups"] = ", ".join(sorted(g for g in a["groups"] if g))
    if q:
        items = [a for a in items if q in a["item"].lower() or q in a["groups"].lower()]
    return render(request, "recon/budget_mapping.html", {"items": items, "groups": groups(cats), "show": show, "q": q, "counts": counts})


@login_required
@require_POST
def budget_automap(request):
    n = automap_unmapped()
    messages.success(request, f"{n} budget line{'s were' if n != 1 else ' was'} matched to a report line." if n else
                              "Nothing new could be matched automatically. Choose the report line yourself for the rest.")
    return redirect("budget_mapping")
